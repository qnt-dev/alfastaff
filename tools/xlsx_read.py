"""Чтение таблицы «Вакансии ГК Альфа-Стафф.xlsx»: доступные объекты (видимые листы и видимые колонки)
и значения по строкам-подписям (колонка A). Используется скриптом build_objects.py."""
import re, openpyxl
from openpyxl.utils import column_index_from_string as CI

SKIP_SHEETS = {'Минск', 'Адреса офисов'}           # Беларусь не показываем; адреса офисов — отдельно
SKIP_NAMES = re.compile(r'кухонный цех|^хостел$', re.I)   # вакансии внутри компании, не объекты
STOP = re.compile(r'стоп|приостанов|набора нет|нет набора', re.I)

def clean(v):
    if v is None: return ''
    s = str(v).replace('\xa0', ' ').replace('\r', '')
    s = '\n'.join(re.sub(r'[ \t]+', ' ', l).strip() for l in s.split('\n'))
    return s.strip()

def read(path):
    wb = openpyxl.load_workbook(path)
    out = []
    for ws in wb.worksheets:
        if ws.sheet_state != 'visible' or ws.title in SKIP_SHEETS: continue
        hidden = set()
        for key, d in ws.column_dimensions.items():
            if d.hidden:
                a = d.min or CI(key); b = d.max or a
                hidden.update(range(a, b + 1))
        grid = {}
        for row in ws.iter_rows():
            for c in row:
                if c.value is not None: grid[(c.row, c.column)] = c.value
        for mr in ws.merged_cells.ranges:
            v = grid.get((mr.min_row, mr.min_col))
            if v is None: continue
            for r in range(mr.min_row, mr.max_row + 1):
                for cc in range(mr.min_col, mr.max_col + 1): grid[(r, cc)] = v
        # строки-подписи: начало каждой — непустая ячейка в колонке A (у объединений — верхняя строка)
        labels = []
        for r in range(2, ws.max_row + 1):
            a = ws.cell(r, 1).value
            if a is not None and str(a).strip():
                labels.append((r, clean(a).split('\n')[0]))
        bounds = [(r, labels[i + 1][0] - 1 if i + 1 < len(labels) else min(ws.max_row, r + 12), lab) for i, (r, lab) in enumerate(labels)]
        # подписи второго уровня (колонка B: «медицинская книжка», «спец.одежда»)
        sub = {r: clean(ws.cell(r, 2).value) for r in range(2, ws.max_row + 1) if ws.cell(r, 2).value}
        heads = []
        for mr in ws.merged_cells.ranges:
            if mr.min_row == 1 and mr.min_col >= 3: heads.append((mr.min_col, mr.max_col))
        for c in range(3, ws.max_column + 1):
            if ws.cell(1, c).value and not any(a <= c <= b for a, b in heads): heads.append((c, c))
        heads.sort()
        for a, b in heads:
            cols = [c for c in range(a, b + 1) if c not in hidden]
            if not cols: continue
            name = clean(grid.get((1, a)))
            if not name or SKIP_NAMES.search(name): continue
            fields = {}
            for r0, r1, lab in bounds:
                key = lab
                rows = []
                for r in range(r0, r1 + 1):
                    vals = []
                    for c in cols:
                        v = clean(grid.get((r, c)))
                        if v and v not in vals: vals.append(v)
                    if vals:
                        k2 = sub.get(r)
                        rows.append(((k2 + ': ') if k2 and k2 != lab else '') + ' | '.join(vals))
                if rows: fields[key] = rows
            out.append({'sheet': ws.title, 'name': name, 'stop': bool(STOP.search(name)), 'fields': fields})
    return out

def offices(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb['Адреса офисов']; res = []
    for r in range(2, ws.max_row + 1):
        a = clean(ws.cell(r, 1).value)
        if a and ws.cell(r, 3).value:
            res.append({'address': a, 'how': clean(ws.cell(r, 2).value), 'hours': clean(ws.cell(r, 3).value), 'map': clean(ws.cell(r, 4).value), 'city': clean(ws.cell(r - 1, 1).value)})
    return res
